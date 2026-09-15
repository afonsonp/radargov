# -*- coding: utf-8 -*-
"""O registo da casa: o que a empresa fez com cada concurso.

Primeira fonte: o Excel de analise de concursos do Afonso
(Analise_Concursos_Publicos.xlsm), exportado de uma lista do SharePoint
e completado a mao numa folha por concurso. Le-se com a MESMA logica dos
consolidadores VBA que ele tem la dentro -- as folhas C_ sao os registos,
as tabelas planas sao derivadas e podem estar desactualizadas.

Cada linha liga-se ao PROCEDIMENTO do radar (o anuncio original, nunca
uma alteracao) e, quando o estado e inequivoco, escreve-se a triagem e o
quadro nesse anuncio. O que o Excel sabe e o radar nao (concorrentes,
precos por perfil, EBITDA, perfis exigidos) fica na tabela `casa` e
mostra-se na ficha.

E o primeiro modulo fora do radar.py (decisao de 01/09/2026): importa o
radar de dentro das funcoes, porque o radar importa este para as rotas.
"""
import json
import re
import time
from datetime import datetime

# Entidades que nao entram (decisao do Afonso a 01/09/2026): as espanholas.
# O universo do radar e a parte L do DR, por isso nada disto se ligaria
# de qualquer maneira; a regra existe para nao ficarem "por ligar" para
# sempre a pedir atencao.
FORA_DO_PAIS = ("asturias", "principado de", "xunta de", "junta de andalucia",
                "generalitat", "ayuntamiento", "gobierno de", "ministerio de",
                "comunidad de", "diputacion")

# Os estados do Excel que se traduzem em triagem do radar. "Cancelado" e
# "TBD" ficam so no registo: nao ha estado do radar que os diga sem mentir.
ESTADOS_COM_TRIAGEM = ("nao fomos", "submetido", "perdido", "ganho")

# As fases do Zoho, no vocabulario da casa. "2.3 - Negotiation" e
# "Ready for Proposal" nao sao um fim: a primeira ja concorremos, a
# segunda ainda nem propusemos -- e o Excel so tem "Submetido" para o
# meio do caminho, por isso e ai que ambas caem. Uma fase que nao esteja
# aqui nao traduz, e a linha fica com o que o Excel diz.
# As chaves sao o que o `_norma()` devolve -- que guarda os pontos e os
# hifens, e por isso "2.3 - negotiation" fica assim mesmo. Ha teste.
TRADUCAO_ZOHO = {"lost": "Perdido", "won": "Ganho", "cancel": "Cancelado",
                 "2.3 - negotiation": "Submetido",
                 "ready for proposal": "Submetido"}

# As razoes de nao participacao do Excel, no vocabulario dos motivos de
# abandono (MOTIVOS_ABANDONO). Chaves ja simplificadas.
MAPA_RAZAO = {"preco base demasiado baixo": "Preço base baixo",
              "preco base baixo": "Preço base baixo",
              "falta de certificacoes": "Falta de certificações",
              "falta de cv s": "Falta de CV's",
              "falta de cvs": "Falta de CV's",
              "fora do nosso ambito": "Fora do âmbito",
              "fora do ambito": "Fora do âmbito",
              "prazo de entrega curto": "Prazo curto",
              "prazo curto": "Prazo curto"}

STOP = set("de da do das dos e a o as os em para com no na nos nas por um uma "
           "ao aos servicos servico aquisicao contratacao prestacao fornecimento "
           "projeto projecto desenvolvimento".split())

# Nomes curtos que o Excel usa e que o corpus nao resolve sozinho.
ALIAS = {"ipl": "instituto politecnico de leiria",
         "tml": "transportes metropolitanos de lisboa",
         "osae": "ordem dos solicitadores",
         "act": "autoridade para as condicoes do trabalho",
         "icnf": "conservacao da natureza",
         "igefe": "gestao financeira e equipamentos",
         "igfej": "gestao financeira e equipamentos da justica",
         "adene": "agencia para a energia",
         "ifap": "financiamento da agricultura",
         "impic": "mercados publicos do imobiliario",
         "aicep": "aicep",
         "ipdj": "portugues do desporto e juventude",
         "dgpj": "direcao geral da politica de justica",
         "ipca": "instituto politecnico do cavado",
         "ama": "modernizacao administrativa",
         "ansr": "seguranca rodoviaria",
         "inem": "emergencia medica",
         "lneg": "energia e geologia",
         "dgt": "direcao geral do territorio",
         "u porto": "universidade do porto",
         "emrp": "emrp"}

LIMIAR = 0.6           # abaixo disto nao ha ligacao
FOLGA = 0.15           # o segundo tem de ficar a esta distancia do primeiro
MAX_CANDIDATOS = 5


# ------------------------------------------------------------ o Excel

def _num(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    try:
        import radar
        return radar.euros_do_texto(str(v)) or None
    except Exception:
        return None


def _txt(v):
    return " ".join(str(v).split()) if v is not None else ""


def _celulas(ws):
    """A folha como lista de linhas (listas), sem depender de dimensions."""
    return [list(r) for r in ws.iter_rows(values_only=True)]


def _procura(linhas, texto, coluna=0, a_partir=0):
    """Indice da primeira linha cuja coluna comeca por `texto` (maiusc.)."""
    alvo = texto.upper()
    for i in range(a_partir, len(linhas)):
        cel = linhas[i][coluna] if coluna < len(linhas[i]) else None
        if cel is not None and _txt(cel).upper().startswith(alvo):
            return i
    return -1


def _valor_em(linhas, i, coluna=1):
    if 0 <= i < len(linhas) and coluna < len(linhas[i]):
        return linhas[i][coluna]
    return None


def ler_folha_concurso(ws):
    """Uma folha C_: as tabelas A, A.1, B, C e C.1, pelas mesmas ancoras
    que as macros usam ("TABELA B", "TABELA C", cabecalhos PERFIL e
    CONCORRENTE)."""
    L = _celulas(ws)
    if not L:
        return None
    fora = {"entidade": _txt(_valor_em(L, 0, 0)),
            "nome": _txt(_valor_em(L, 0, 1)),
            "id": None, "perfis": [], "concorrentes": [], "precos_perfis": []}
    # o id estavel fica em Z1 (coluna 26, indice 25) com a marca em Z2
    z1 = _valor_em(L, 0, 25)
    if isinstance(z1, (int, float)):
        fora["id"] = int(z1)
    else:
        try:
            fora["id"] = int(_txt(_valor_em(L, 0, 2)) or 0) or None
        except ValueError:
            pass
    # Tabela A: rotulo na coluna A, valor na B
    ia = _procura(L, "TABELA A")
    rotulos = {"modelo": "modelo", "prazo de execu": "prazo_meses",
               "preço base": "preco_base", "preco base": "preco_base",
               "critério": "criterio", "criterio": "criterio",
               "plataforma": "plataforma", "ano": "ano", "status": "status"}
    if ia >= 0:
        for i in range(ia + 1, min(ia + 12, len(L))):
            rot = _txt(_valor_em(L, i, 0)).lower()
            for chave, campo in rotulos.items():
                if rot.startswith(chave):
                    fora[campo] = _valor_em(L, i, 1)
    # A.1: perfis exigidos, ate a linha vazia ou a TABELA B
    ip = _procura(L, "PERFIL", 0, ia if ia >= 0 else 0)
    if ip >= 0:
        for i in range(ip + 1, len(L)):
            p = _txt(_valor_em(L, i, 0))
            if not p or p.upper().startswith("TABELA"):
                break
            fora["perfis"].append({
                "perfil": p, "tecnologias": _txt(_valor_em(L, i, 1)),
                "anos": _num(_valor_em(L, i, 2)), "n": _num(_valor_em(L, i, 3)),
                "horas": _num(_valor_em(L, i, 4)),
                "certificacoes": _txt(_valor_em(L, i, 5))})
    # Tabela B: nove valores na coluna B, pela ordem do modelo
    ib = _procura(L, "TABELA B")
    if ib >= 0:
        fora["valor_proposta"] = _num(_valor_em(L, ib + 1))
        fora["lugar"] = _num(_valor_em(L, ib + 2))
        fora["ebitda"] = _num(_valor_em(L, ib + 3))
        fora["gap_base"] = _num(_valor_em(L, ib + 4))
        fora["gap_base_pct"] = _num(_valor_em(L, ib + 5))
        fora["gap_primeiro"] = _num(_valor_em(L, ib + 6))
        fora["gap_primeiro_pct"] = _num(_valor_em(L, ib + 7))
        fora["razao"] = _txt(_valor_em(L, ib + 8))
        fora["notas"] = _txt(_valor_em(L, ib + 9))
    # Tabela C: cinco lugares, tres linhas cada (titulo, nome, valor)
    ic = _procura(L, "TABELA C")
    if ic >= 0:
        for lugar in range(1, 6):
            base = ic + 1 + (lugar - 1) * 3
            nome = _txt(_valor_em(L, base + 1))
            if nome:
                fora["concorrentes"].append(
                    {"lugar": lugar, "nome": nome,
                     "valor": _num(_valor_em(L, base + 2))})
    # C.1: preco por perfil por concorrente
    icc = _procura(L, "CONCORRENTE", 0, ic if ic >= 0 else 0)
    if icc >= 0:
        for i in range(icc + 1, len(L)):
            conc = _txt(_valor_em(L, i, 0))
            if not conc or conc.upper().startswith("LEGENDA"):
                break
            fora["precos_perfis"].append({
                "concorrente": conc, "perfil": _txt(_valor_em(L, i, 1)),
                "tecnologias": _txt(_valor_em(L, i, 2)),
                "anos": _num(_valor_em(L, i, 3)), "valor": _num(_valor_em(L, i, 4)),
                "horas": _num(_valor_em(L, i, 5)), "hora": _num(_valor_em(L, i, 6))})
    return fora


def ler_excel(caminho):
    """As linhas do INDICE, completadas pela folha de cada concurso quando
    existe. Devolve uma lista de dicionarios, um por concurso, com `id`
    (a coluna K do INDICE, estavel) e `folha` (o nome da folha ou '')."""
    import openpyxl
    wb = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
    if "ÍNDICE" not in wb.sheetnames:
        raise ValueError("o ficheiro não tem a folha ÍNDICE")
    L = _celulas(wb["ÍNDICE"])
    cab = [_txt(c).upper() for c in (L[0] if L else [])]
    def col(nome, omissao):
        for i, c in enumerate(cab):
            if c.startswith(nome):
                return i
        return omissao
    ci = {"nome": col("NOME", 0), "entidade": col("ENTIDADE", 1),
          "modelo": col("MODELO", 2), "prazo": col("PRAZO", 3),
          "preco": col("PRE", 4), "criterio": col("CRIT", 5),
          "plataforma": col("PLATAFORMA", 6), "ano": col("ANO", 7),
          "status": col("STATUS", 8), "id": col("ID", 10)}
    folhas = {}
    for nome in wb.sheetnames:
        if nome.startswith("C_"):
            f = ler_folha_concurso(wb[nome])
            if f and f["id"]:
                f["folha"] = nome
                folhas[f["id"]] = f
    linhas = []
    for r in L[1:]:
        if not r or not _txt(r[ci["nome"]] if ci["nome"] < len(r) else None):
            continue
        def g(k):
            i = ci[k]
            return r[i] if i < len(r) else None
        try:
            ide = int(g("id")) if g("id") not in (None, "") else None
        except (TypeError, ValueError):
            ide = None
        linha = {"id": ide, "nome": _txt(g("nome")), "entidade": _txt(g("entidade")),
                 "modelo": _txt(g("modelo")), "prazo_meses": _num(g("prazo")),
                 "preco_base": _num(g("preco")), "criterio": _txt(g("criterio")),
                 "plataforma": _txt(g("plataforma")),
                 "ano": int(_num(g("ano")) or 0) or None,
                 "status": _txt(g("status")), "folha": "",
                 "razao": "", "valor_proposta": None, "lugar": None,
                 "ebitda": None, "notas": "", "perfis": [], "concorrentes": [],
                 "precos_perfis": []}
        f = folhas.get(ide) if ide else None
        if f:
            for k in ("razao", "valor_proposta", "lugar", "ebitda", "notas",
                      "perfis", "concorrentes", "precos_perfis", "folha"):
                if f.get(k) not in (None, "", []):
                    linha[k] = f[k]
            # a folha e a versao mais completa da Tabela A: se o INDICE
            # estiver vazio num campo e a folha nao, vale a folha
            for k in ("modelo", "criterio", "plataforma", "status"):
                if not linha[k] and _txt(f.get(k)):
                    linha[k] = _txt(f.get(k))
            if linha["preco_base"] is None and _num(f.get("preco_base")):
                linha["preco_base"] = _num(f.get("preco_base"))
        linhas.append(linha)
    wb.close()
    return linhas


def fora_do_pais(linha):
    import radar
    texto = radar.simplifica(" ".join((linha.get("entidade") or "",
                                       linha.get("notas") or "")))
    return any(p in texto for p in FORA_DO_PAIS)


# --------------------------------------------------------- a ligacao

def _toks(texto):
    import radar
    return {w for w in radar.simplifica(texto or "").split()
            if w not in STOP and len(w) > 2}


def _norma(texto):
    import radar
    return radar.simplifica(texto or "")


class Acervo(object):
    """Os anuncios candidatos, lidos UMA vez por importacao: so os
    originais (nunca alteracoes), so o DR, dentro dos anos pedidos."""

    def __init__(self, c, anos):
        anos = sorted(a for a in anos if a)
        if not anos:
            anos = [datetime.now().year]
        de, ate = "%d-01-01" % (min(anos) - 1), "%d-12-31" % max(anos)
        self.linhas = []
        for r in c.execute(
                "SELECT ref, titulo_norm, entidade_norm, nif, preco_base, "
                "data_pub, estado FROM anuncios WHERE data_pub BETWEEN ? AND ? "
                "AND estado != 'alteracao' AND COALESCE(fonte,'dr')='dr'",
                (de, ate)):
            self.linhas.append((r["ref"],
                                {w for w in (r["titulo_norm"] or "").split()
                                 if w not in STOP and len(w) > 2},
                                r["entidade_norm"] or "", r["nif"] or "",
                                _num(r["preco_base"]), r["data_pub"] or ""))
        self.por_ref = {l[0]: l for l in self.linhas}

    def actualiza(self, c, ref):
        """Depois de se ler o detalhe de um candidato: o preco base passa
        a contar, e se o detalhe o revelou como alteracao de outro (o
        _guardar_detalhe ja o ligou ao original) sai da lista -- era a
        razao de muitas ambiguidades: republicacoes por ler."""
        r = c.execute("SELECT ref, titulo_norm, entidade_norm, nif, preco_base, "
                      "data_pub, estado FROM anuncios WHERE ref=?", (ref,)).fetchone()
        if not r or r["estado"] == "alteracao":
            self.por_ref.pop(ref, None)
            self.linhas = [l for l in self.linhas if l[0] != ref]
            return
        nova = (r["ref"], {w for w in (r["titulo_norm"] or "").split()
                           if w not in STOP and len(w) > 2},
                r["entidade_norm"] or "", r["nif"] or "",
                _num(r["preco_base"]), r["data_pub"] or "")
        self.por_ref[ref] = nova
        self.linhas = [nova if l[0] == ref else l for l in self.linhas]


# Palavras de nome de entidade que nao distinguem ninguem.
STOP_ENT = set("servicos instituto municipio direcao geral regional portugal "
               "nacional autoridade agencia unidade local saude entidade publica "
               "administracao camara municipal secretaria ministerio".split())

_ENTIDADES = {}      # nome no Excel -> (chave no corpus, nome canonico norm., nome norm.)


def _entidade(nome):
    """(chave, canonico, norma) da entidade como o Excel a escreve. A chave
    e o NIF no corpus (quando o nome la esta: SPMS, eSPap, INCM...); o
    canonico e o nome mais usado por essa chave, que e o que os anuncios
    do DR escrevem por extenso. Sem corpus fica so o nome, com alias."""
    import radar
    if nome in _ENTIDADES:
        return _ENTIDADES[nome]
    norma = _norma(nome)
    norma = ALIAS.get(norma, norma)
    chave, canon = "", ""
    try:
        chave = radar.entidade_do_anuncio("", nome or "")
        if chave and radar.ha_corpus():
            with radar.liga_corpus() as k:
                r = k.execute("SELECT nome FROM entidades WHERE chave=?",
                              (chave,)).fetchone()
            canon = radar.simplifica(r["nome"]) if r else ""
    except Exception:           # sem corpus (contratos.db) nao ha chave; fica o nome
        chave, canon = "", ""
    _ENTIDADES[nome] = (chave, canon, norma)
    return _ENTIDADES[nome]


def pontuar(linha, acervo):
    """[(pontuacao, ref)] por ordem decrescente, para uma linha do Excel.

    O nome no Excel e uma abreviatura do titulo do DR ("Plataforma
    Central de Deteção Precoce" para "(DAG) Aquisição de serviços para
    evolução da Plataforma Central de Deteção Precoce no âmbito..."):
    conta sobretudo a CONTENCAO -- que fraccao das palavras do Excel esta
    no titulo --, e o Jaccard so desempata. Medido a 02/09/2026: com o
    Jaccard sozinho ficavam 105 das 187 linhas ambiguas."""
    T = _toks(linha["nome"])
    if not T:
        return []
    chave, canon, ent = _entidade(linha["entidade"])
    palavras = [w for w in (canon or ent).split() if len(w) > 4 and w not in STOP_ENT]
    ano = str(linha["ano"]) if linha["ano"] else ""
    anos = {ano, str(linha["ano"] - 1)} if ano else None
    fora = []
    for ref, TT, EE, nif, pb, dp in acervo.linhas:
        if anos and dp[:4] not in anos:
            continue
        inter = len(T & TT)
        if not inter:
            continue
        sc = 0.7 * inter / float(len(T)) + 0.3 * inter / float(len(T | TT))
        if ano and dp[:4] == ano:
            sc += 0.05          # o ano do Excel e o da decisao; quase sempre o do DR
        ent_ok = ((chave and nif and chave == nif)
                  or (canon and (canon in EE or (EE and EE in canon)))
                  or (ent and ent in EE)
                  or (palavras and sum(w in EE for w in palavras)
                      >= max(1, len(palavras) // 2)))
        if ent_ok:
            sc += 0.35
        if linha["preco_base"] and pb and abs(pb - linha["preco_base"]) < 0.5:
            sc += 0.5
        fora.append((round(sc, 3), ref))
    fora.sort(reverse=True)
    return fora


def ref_pelo_base(c, linha, acervo, pontos=()):
    """O anuncio pelo contrato celebrado. O valor do 1.º lugar da tabela C
    (ou a proposta da casa, quando ganhou) e o preco contratual no BASE, e
    o BASE guarda o `n_anuncio`, que e o ref do DR. Medido a 01/09/2026:
    23 dos 29 primeiros lugares com valor acham-se assim. Um valor
    sozinho engana (o mesmo numero em suturas e em software): exige-se
    a entidade certa ou, sem chave, algum parentesco no titulo."""
    import radar
    if not radar.ha_corpus():
        return ""
    st = _norma(linha.get("status"))
    valores = [cc["valor"] for cc in linha.get("concorrentes") or []
               if cc.get("lugar") == 1 and cc.get("valor")]
    if st == "ganho" and linha.get("valor_proposta"):
        valores.append(linha["valor_proposta"])
    if not valores:
        return ""
    chave, _, _ = _entidade(linha["entidade"])
    ano = linha.get("ano") or datetime.now().year
    achados = {}
    with radar.liga_corpus() as k:
        for v in valores:
            for r in k.execute(
                    "SELECT n_anuncio, adjudicante_chave FROM contratos "
                    "WHERE preco_contratual BETWEEN ? AND ? AND data_celebracao >= ? "
                    "AND n_anuncio != ''", (v - 1, v + 1, "%d-01-01" % (ano - 1))):
                achados[r["n_anuncio"]] = r["adjudicante_chave"] or ""
    score = dict((r, s) for s, r in pontos)
    refs = set()
    for n, adj in achados.items():
        if chave and adj and adj != chave:
            continue
        ref = n
        if n not in acervo.por_ref:
            a = c.execute("SELECT estado, altera FROM anuncios WHERE ref=?",
                          (n,)).fetchone()
            if not a:
                continue
            if a["estado"] == "alteracao":
                ref = radar.raiz_da_alteracao(c, n, a["altera"] or "") or ""
            if ref not in acervo.por_ref:
                continue
        if (chave and adj == chave) or score.get(ref, 0) >= 0.3:
            refs.add(ref)
    return refs.pop() if len(refs) == 1 else ""


def decidir(pontos):
    """(ref, candidatos): a ligacao unica, ou '' e a lista dos candidatos."""
    if not pontos or pontos[0][0] < LIMIAR:
        return "", [r for s, r in pontos[:MAX_CANDIDATOS] if s >= 0.3]
    if len(pontos) > 1 and pontos[1][0] >= LIMIAR and pontos[1][0] > pontos[0][0] - FOLGA:
        return "", [r for s, r in pontos[:MAX_CANDIDATOS] if s >= 0.3]
    return pontos[0][1], []


# ------------------------------------------------------- a aplicacao

def _texto_top3(concorrentes):
    import radar
    partes = []
    for cc in sorted(concorrentes or [], key=lambda x: x.get("lugar") or 9)[:3]:
        v = cc.get("valor")
        partes.append("%d.º %s%s" % (cc.get("lugar") or 0, cc.get("nome") or "",
                                     " " + radar._texto_do_preco(v) if v else ""))
    return " · ".join(partes)[:300]


def estado_efectivo(linha):
    """O estado que vale, entre o que o Excel diz e o que o Zoho diz.

    A regra e dele, dada a 03/09/2026 depois de ver os numeros: **o
    "Nao fomos" do Excel prevalece, e e o unico**; em tudo o resto ganha
    o Zoho, que e a fonte mais actual. A excepcao existe porque o Zoho
    nao tem palavra para "nao concorremos" -- em 46 das 92 linhas que
    cruzam, o Excel diz "Nao fomos" e o Zoho diz "Lost". Sem a excepcao,
    metade do cruzamento perdia a distincao.

    **E o Zoho tambem nao decide uma linha que e UM LOTE.** As duas
    fontes contam coisas diferentes: o Excel tem uma linha por lote, o
    Zoho um negocio por procedimento. Um "Won" do Zoho quer dizer
    "ganhamos pelo menos um lote" e nao diz nada sobre este. Medido no
    1947/2026 (04/09/2026, e a razao desta regra existir): tres lotes,
    tres linhas -- #14 o L1 perdido, #97 o L2 ganho, #98 o L3 perdido --
    e no Zoho um so negocio, "Won". Sem esta guarda, o #14 passava de
    Perdido a Ganho. Uma linha de lote fica com o que o Excel diz, que e
    a fonte fina para lotes; o `lote = 0` (o conjunto) nao e um lote e
    aceita o Zoho como qualquer outra.

    Nao le a base e nao escreve nada: e derivada, de proposito. O
    `status` continua a ser o do Excel e o `zoho_fase` o do Zoho, cada
    um intacto na sua coluna -- assim uma reimportacao do Excel nao
    desfaz a regra, e mudar a regra nao obriga a reescrever dados.
    """
    st = linha.get("status")
    if _norma(st) == "nao fomos":
        return st
    if linha.get("lote"):           # >= 1: o zero e o conjunto, e nao conta
        return st
    return TRADUCAO_ZOHO.get(_norma(linha.get("zoho_fase")), st)


def estado_pretendido(linha):
    """(estado da escada, campos) que o registo da casa pede, ou None
    quando o estado do Excel nao se traduz em nada.

    O estado vem do `estado_efectivo()`, nao do `status` cru: quem manda
    e o Zoho, tirando o "Nao fomos". Desde 15/09/2026 devolve uma das
    oito palavras da casa (radar.ESTADOS_DA_CASA) em vez de um par
    estado+fase: o vocabulario passou a ser um so, e "interessa" com uma
    fase ao lado era o mesmo estado dito duas vezes.
    """
    import radar
    st = _norma(estado_efectivo(linha))
    if st not in ESTADOS_COM_TRIAGEM:
        return None
    if st == "nao fomos":
        # a razao ja canonica (vem do modelo, ou e uma das do radar) fica
        # como esta; o mapa e para as variantes do Excel antigo
        razao = (linha.get("razao") or "").strip()
        motivo = MAPA_RAZAO.get(_norma(razao)) or razao or None
        return ("nao_fomos", {"motivo": motivo})
    campos = {}
    if linha.get("valor_proposta"):
        campos["valor_proposta"] = radar._texto_do_preco(linha["valor_proposta"])
    if st == "submetido":
        return ("submetido", campos)
    campos["lugar"] = (int(linha["lugar"]) if linha.get("lugar")
                       else (1 if st == "ganho" else None))
    campos["top3"] = _texto_top3(linha.get("concorrentes")) or None
    return (st, campos)


def aplicar(c, linha, ref, quem="Excel"):
    """Escreve o que o registo da casa sabe na PROPOSTA do anuncio ligado.
    Devolve 'aplicado', 'igual', 'sem estado', 'conflito' ou 'sem anúncio'.

    Ate 15/09/2026 escrevia no `anuncios` (estado, fase_id e as colunas do
    quadro). Agora cria ou move uma proposta -- e por decisao dele nesse
    dia (D4 do docs/historico/CRM.md) o Excel deixou de ser fonte
    permanente: serve para trazer os concursos passados e o resultado
    deles. **Nao passa por cima de uma decisao humana feita no radar**:
    uma proposta que ja esteja noutra ranhura fica como esta, e o
    conflito e registado uma vez no historico.

    O LOTE vem da linha do Excel: o Excel tem uma linha por lote e a
    escada tem uma proposta por lote, o que finalmente e a mesma coisa.
    """
    import radar
    if not c.execute("SELECT 1 FROM anuncios WHERE ref=?", (ref,)).fetchone():
        return "sem anúncio"
    pedido = estado_pretendido(linha)
    if not pedido:
        return "sem estado"
    estado, campos = pedido
    lote = linha.get("lote")
    p = c.execute("SELECT * FROM propostas WHERE ref=? AND "
                  "COALESCE(lote,-1)=COALESCE(?,-1)", (ref, lote)).fetchone()
    if p:
        igual = (p["estado"] == estado
                 and all((p[k] or None) == (v or None)
                         for k, v in campos.items() if k in p.keys()))
        if igual:
            return "igual"
        if p["estado"] != estado:
            aviso = ("o registo da casa diz «%s»; mantém-se a decisão do radar"
                     % (linha.get("status") or ""))
            if not c.execute("SELECT 1 FROM historico WHERE ref=? AND detalhe=?",
                             (ref, aviso)).fetchone():
                _registar(c, ref, "estado", aviso, quem)
            return "conflito"
        id_ = p["id"]
    else:
        # Criada aqui e nao pelo radar.criar_proposta(): esta funcao corre
        # DENTRO da transaccao da importacao, com a ligacao `c` aberta, e
        # abrir uma segunda ligacao a meio trancava a base -- e um erro
        # que este modulo ja pagou uma vez.
        a = c.execute("SELECT titulo, entidade, preco_base, lotes FROM anuncios "
                      "WHERE ref=?", (ref,)).fetchone()
        cur = c.execute(
            "INSERT INTO propostas (ref, lote, entidade, titulo, estado, "
            "preco_base, criada_em) VALUES (?,?,?,?,?,?,?)",
            (ref, lote, a["entidade"] or "", a["titulo"] or "", estado,
             radar.preco_base_do_lote(a, lote),
             datetime.now().strftime("%Y-%m-%d %H:%M")))
        id_ = cur.lastrowid
    sets, vals = ["estado=?"], [estado]
    # O carimbo que faz o funil esvaziar, e a mesma regra do radar: so as
    # ranhuras fechadas o levam, e sair delas limpa-o.
    sets.append("fechada_em=?")
    vals.append(datetime.now().strftime("%Y-%m-%d %H:%M")
                if estado in radar.ESTADOS_FECHADOS else None)
    if "motivo" not in campos:
        sets.append("motivo=NULL")
    for k, v in campos.items():
        sets.append("%s=?" % k)
        vals.append(v)
    c.execute("UPDATE propostas SET %s WHERE id=?" % ", ".join(sets), vals + [id_])
    detalhe = ("%s%s, do registo da casa"
               % (radar.estado_da_casa(estado),
                  " (%s)" % campos["motivo"] if campos.get("motivo") else ""))
    _registar(c, ref, "estado", detalhe, quem)
    if campos.get("valor_proposta"):
        _registar(c, ref, "preço proposto", campos["valor_proposta"], quem)
    if campos.get("lugar") or campos.get("top3"):
        _registar(c, ref, "relatório preliminar",
                  "%s%s" % ("%dº lugar" % campos["lugar"] if campos.get("lugar")
                            else "sem lugar",
                            " — " + campos["top3"] if campos.get("top3") else ""),
                  quem)
    return "aplicado"


def _registar(c, ref, accao, detalhe, quem):
    c.execute("INSERT INTO historico (ref, quem, accao, detalhe, quando) "
              "VALUES (?,?,?,?,?)",
              (ref, quem, accao, detalhe, datetime.now().strftime("%Y-%m-%d %H:%M")))


# --------------------------------------------------------- a tabela

def iniciar_tabelas(c):
    c.execute("""CREATE TABLE IF NOT EXISTS casa (
        id INTEGER PRIMARY KEY, nome TEXT, entidade TEXT, modelo TEXT,
        prazo_meses REAL, preco_base REAL, criterio TEXT, plataforma TEXT,
        ano INTEGER, status TEXT, razao TEXT, valor_proposta REAL,
        lugar INTEGER, ebitda REAL, notas TEXT, folha TEXT,
        perfis TEXT, concorrentes TEXT, precos_perfis TEXT,
        ref TEXT, ligacao TEXT DEFAULT '', candidatos TEXT DEFAULT '[]',
        fora INTEGER DEFAULT 0, resultado TEXT DEFAULT '',
        importado_em TEXT, aplicado_em TEXT)""")
    c.execute("CREATE INDEX IF NOT EXISTS ix_casa_ref ON casa(ref)")
    # Porque e que uma linha nao tem anuncio: "consulta previa", "antes de
    # 2025", "nao sei". Vem das respostas do Afonso (02/09/2026) e e o
    # que distingue "por ligar" de "nao ha nada para ligar".
    colunas = [r["name"] for r in c.execute("PRAGMA table_info(casa)")]
    if "porque_sem_ref" not in colunas:
        c.execute("ALTER TABLE casa ADD COLUMN porque_sem_ref TEXT")
    # A que lote do anuncio esta linha corresponde (o Excel tem uma linha
    # por lote; o DR um anuncio para todos). NULL = anuncio sem lotes, ou
    # lote por identificar. ZERO = o conjunto: a linha e do procedimento
    # inteiro, nao de um lote (resposta do Afonso a 03/09/2026 sobre as
    # linhas #23 e #26, cujo preco e o total do anuncio).
    if "lote" not in colunas:
        c.execute("ALTER TABLE casa ADD COLUMN lote INTEGER")
    # O que o Zoho diz do mesmo concurso, em coluna PROPRIA -- nao por
    # cima do `status`, e de proposito. Medido a 03/09/2026 no
    # cruzamento das 148 oportunidades da vista dos Negocios: em 46 das
    # 92 linhas que cruzam, o Excel diz "Nao fomos" e o Zoho diz "Lost".
    # O Zoho nao tem palavra para "nao concorremos", e escrever por cima
    # apagava a distincao. Fica cada um com a sua coluna, e quem manda
    # decide-se quando o vocabulario dos estados estiver decidido.
    # `zoho_como` guarda por que regra a linha casou, para a ligacao ser
    # auditavel: um cruzamento por semelhanca de nome nao e uma certeza.
    for coluna, tipo in (("zoho_fase", "TEXT"), ("zoho_montante", "REAL"),
                         ("zoho_como", "TEXT"), ("zoho_em", "TEXT")):
        if coluna not in colunas:
            c.execute("ALTER TABLE casa ADD COLUMN %s %s" % (coluna, tipo))


RX_LOTE_NO_NOME = re.compile(r"\bL(?:ote)?\s*\.?\s*(\d{1,2})\b", re.I)


def lote_da_linha(linha, lotes):
    """O numero do lote a que a linha do Excel corresponde, 0 ou None.

    Primeiro pelo preco base: a linha traz o preco base DO LOTE (medido
    nos quatro anuncios com varias linhas a 02/09/2026 -- #7 = 53 667,20
    = lote 1 do 2770/2026). Depois por um "L1" ou "Lote 2" no nome.

    Em ultimo, ZERO -- o conjunto -- quando o preco da linha e a SOMA de
    todos os lotes: nesse caso o numero do Excel e o total do anuncio e a
    linha nao esta dividida por lotes (o Afonso, a 03/09/2026, sobre as
    #23 e #26). Zero e falso em Python, e e de proposito: quem contava
    "linhas com lote" continua a nao as contar, mas deixa de as confundir
    com as que estao mesmo por identificar."""
    if not lotes:
        return None
    pb = linha.get("preco_base")
    if pb:
        for l in lotes:
            v = _num(l.get("preco_base"))
            if v and abs(v - pb) < 1:
                return l["n"]
    m = RX_LOTE_NO_NOME.search(linha.get("nome") or "")
    if m and 1 <= int(m.group(1)) <= max(l["n"] for l in lotes):
        return int(m.group(1))
    if pb:
        soma = [_num(l.get("preco_base")) for l in lotes]
        if all(soma) and abs(sum(soma) - pb) < 1:
            return 0
    return None


ESTADOS_DE_LOTE = ("ganho", "perdido", "submetido", "nao fomos")


def estado_do_lote(linha):
    """O estado de UMA linha da casa, como chave: 'ganho', 'perdido',
    'submetido', 'nao fomos' ou '' quando o Excel nao diz nada de util
    ("Cancelado", "TBD"). E o estado_efectivo() normalizado -- por isso
    uma linha de lote fica com o que o Excel diz, e o conjunto (lote 0)
    aceita o Zoho."""
    st = _norma(estado_efectivo(linha))
    return st if st in ESTADOS_DE_LOTE else ""


def linhas_de_lotes(c, refs):
    """{ref: [linhas da casa com `lote` preenchido]} para varios anuncios
    de uma vez -- o quadro pede pelas suas cartas todas, nao uma a uma."""
    refs = [r for r in refs if r]
    if not refs:
        return {}
    saida = {}
    for i in range(0, len(refs), 400):
        pedaco = refs[i:i + 400]
        for r in c.execute(
                "SELECT id, ref, lote, status, zoho_fase, valor_proposta, lugar, "
                "nome, razao FROM casa WHERE lote IS NOT NULL AND ref IN (%s) "
                "ORDER BY lote, id" % ",".join("?" * len(pedaco)), pedaco):
            saida.setdefault(r["ref"], []).append(dict(r))
    return saida


def lotes_do_anuncio(c, ref):
    r = c.execute("SELECT lotes FROM anuncios WHERE ref=?", (ref,)).fetchone()
    try:
        return json.loads(r["lotes"]) if r and r["lotes"] else []
    except ValueError:
        return []


CAMPOS_EXCEL = ("nome", "entidade", "modelo", "prazo_meses", "preco_base",
                "criterio", "plataforma", "ano", "status", "razao",
                "valor_proposta", "lugar", "ebitda", "notas", "folha")


def _guardar_linha(c, linha, ref, ligacao, candidatos, fora, resultado, agora):
    vals = [linha.get(k) for k in CAMPOS_EXCEL]
    vals = [(None if v == "" else v) for v in vals]
    c.execute(
        "INSERT INTO casa (id, %s, perfis, concorrentes, precos_perfis, ref, "
        "ligacao, candidatos, fora, resultado, importado_em, aplicado_em) "
        "VALUES (?, %s, ?,?,?,?,?,?,?,?,?,?) ON CONFLICT(id) DO UPDATE SET %s, "
        "perfis=excluded.perfis, concorrentes=excluded.concorrentes, "
        "precos_perfis=excluded.precos_perfis, ref=excluded.ref, "
        "ligacao=excluded.ligacao, candidatos=excluded.candidatos, "
        "fora=excluded.fora, resultado=excluded.resultado, "
        "importado_em=excluded.importado_em, aplicado_em=excluded.aplicado_em"
        % (", ".join(CAMPOS_EXCEL), ", ".join("?" * len(CAMPOS_EXCEL)),
           ", ".join("%s=excluded.%s" % (k, k) for k in CAMPOS_EXCEL)),
        [linha["id"]] + vals + [
            json.dumps(linha.get("perfis") or [], ensure_ascii=False),
            json.dumps(linha.get("concorrentes") or [], ensure_ascii=False),
            json.dumps(linha.get("precos_perfis") or [], ensure_ascii=False),
            ref or None, ligacao, json.dumps(candidatos), int(bool(fora)),
            resultado, agora, agora if resultado == "aplicado" else None])


def importar(caminho, ensaio=False, ler=True, quem="Excel", relatar=None,
             triagem=False):
    """Le o Excel, liga cada linha ao procedimento e guarda o registo.

    `ensaio` calcula tudo e nao grava nada. `ler` autoriza ir ao DR
    buscar o detalhe dos candidatos de uma linha ambigua (um pedido por
    candidato, ~1 s) para o preco base desempatar. `triagem` escreve
    tambem a triagem nos anuncios ligados -- DESLIGADO por omissao,
    decisao do Afonso a 02/09/2026: enquanto as ligacoes nao estiverem
    validadas e os lotes (varias linhas do Excel no mesmo anuncio) nao
    tiverem solucao, guarda-se a informacao e mais nada. Devolve o
    relatorio: contagens e as listas do que ficou por ligar."""
    import radar
    diz = relatar or (lambda _: None)
    linhas = ler_excel(caminho)
    linhas = [l for l in linhas if l["id"]]
    agora = datetime.now().strftime("%Y-%m-%d %H:%M")
    rel = {"total": len(linhas), "fora": [], "ligadas": 0, "novas": 0,
           "manuais": 0, "pelo_base": 0, "sem_dr": 0, "em_lotes": 0, "com_lote": 0,
           "conjunto": 0,
           "ambiguas": [], "sem": [],
           "aplicadas": {}, "conflitos": [], "lidos": 0, "ensaio": ensaio}
    _ENTIDADES.clear()
    # Duas passagens, de proposito. A primeira so LE (e vai ao DR pelos
    # candidatos ambiguos); a segunda escreve. Com uma so, a ligacao que
    # ja tinha escrito a primeira linha da casa segurava a base numa
    # transaccao, e o ler_detalhe_de(), que grava pela ligacao dele,
    # ficava a espera ate rebentar com "database is locked" -- so na
    # importacao a serio, porque o ensaio nao escrevia nada.
    decisoes = []
    with radar.liga() as c:
        iniciar_tabelas(c)
    with radar.liga() as c:
        acervo = Acervo(c, {l["ano"] for l in linhas})
        existentes = {r["id"]: dict(r) for r in c.execute(
            "SELECT id, ref, ligacao, porque_sem_ref, zoho_fase, lote "
            "FROM casa")}
        for linha in linhas:
            antes = existentes.get(linha["id"]) or {}
            if fora_do_pais(linha):
                rel["fora"].append(linha["nome"])
                decisoes.append((linha, "", "", [], True))
                continue
            ref, ligacao, candidatos = "", "", []
            if antes.get("ref") and antes.get("ligacao") == "manual":
                ref, ligacao = antes["ref"], "manual"
                rel["manuais"] += 1
            elif antes.get("ligacao") == "nenhum":
                # ele disse que nao ha anuncio no DR: nao se volta a procurar
                ligacao = "nenhum"
                rel["sem_dr"] += 1
            else:
                pontos = pontuar(linha, acervo)
                ref, candidatos = decidir(pontos)
                if not ref:
                    # o contrato celebrado sabe o ref: pelo valor do 1.º lugar
                    ref = ref_pelo_base(c, linha, acervo, pontos)
                    if ref:
                        ligacao, candidatos = "base", []
                        rel["pelo_base"] += 1
                if not ref and candidatos and ler:
                    # desempate pela leitura do detalhe dos candidatos que
                    # ainda nao o tem: traz o preco base, e revela as
                    # republicacoes (que saem da lista). Um pedido por
                    # candidato, ~1 s, sem castigar o portal.
                    for cand in candidatos:
                        if acervo.por_ref.get(cand) and acervo.por_ref[cand][4] is None:
                            ok, _ = radar.ler_detalhe_de(cand)
                            rel["lidos"] += 1
                            time.sleep(0.7)
                            if not ok:
                                break
                            acervo.actualiza(c, cand)
                    ref, candidatos = decidir(pontuar(linha, acervo))
                if ref:
                    ligacao = ligacao or "auto"
                    if antes.get("ref") != ref:
                        rel["novas"] += 1
            decisoes.append((linha, ref, ligacao, candidatos, False))
    with radar.liga() as c:
        for linha, ref, ligacao, candidatos, fora in decisoes:
            # A linha vem do Excel e nao traz nem o que o Zoho diz nem o
            # lote; sem isto, uma reimportacao com --com-triagem
            # desfazia o `estado_efectivo()` em silencio -- e as duas
            # coisas fazem falta, porque o lote e o que TRAVA o Zoho.
            # (O lote desta passagem so se calcula mais a frente, ja
            # depois da triagem; o que interessa aqui e o que esta
            # guardado, que e o mesmo, e existe desde a 1.a importacao.)
            guardado = existentes.get(linha["id"]) or {}
            linha.setdefault("zoho_fase", guardado.get("zoho_fase"))
            linha.setdefault("lote", guardado.get("lote"))
            resultado = ""
            if fora:
                resultado = "fora"
            elif ref:
                rel["ligadas"] += 1
                if not triagem:
                    resultado = "guardado"
                elif ensaio:
                    # o que se faria, sem escrever: le-se o estado actual
                    pedido = estado_pretendido(linha)
                    # O estado a comparar e o da PROPOSTA daquele lote, e
                    # nao o do anuncio: e la que a decisao da casa mora
                    # desde 15/09/2026. Comparar com o do anuncio dava
                    # "aplicado" a tudo, e o ensaio existe precisamente
                    # para dizer o que ia mudar.
                    p = c.execute("SELECT estado FROM propostas WHERE ref=? "
                                  "AND COALESCE(lote,-1)=COALESCE(?,-1)",
                                  (ref, linha.get("lote"))).fetchone()
                    resultado = ("sem estado" if not pedido else
                                 "igual" if p and p["estado"] == pedido[0] else
                                 "aplicado")
                else:
                    resultado = aplicar(c, linha, ref, quem)
                rel["aplicadas"][resultado] = rel["aplicadas"].get(resultado, 0) + 1
                if resultado == "conflito":
                    rel["conflitos"].append((linha["nome"], ref, linha["status"]))
            elif ligacao == "nenhum":
                pass
            elif candidatos:
                rel["ambiguas"].append((linha["id"], linha["nome"], linha["entidade"],
                                        linha["ano"], candidatos))
            else:
                rel["sem"].append((linha["id"], linha["nome"], linha["entidade"],
                                   linha["ano"]))
            if ref:
                lotes = lotes_do_anuncio(c, ref)
                if lotes:
                    rel["em_lotes"] += 1
                    lote = lote_da_linha(linha, lotes)
                    rel["com_lote"] += bool(lote)
                    rel["conjunto"] += lote == 0
                else:
                    lote = None
            if not ensaio:
                _guardar_linha(c, linha, ref, ligacao, candidatos, fora,
                               resultado, agora)
                antes = existentes.get(linha["id"]) or {}
                if antes.get("porque_sem_ref"):
                    c.execute("UPDATE casa SET porque_sem_ref=? WHERE id=?",
                              (antes["porque_sem_ref"], linha["id"]))
                if ref:
                    c.execute("UPDATE casa SET lote=? WHERE id=?", (lote, linha["id"]))
        if not ensaio:
            c.execute("INSERT OR REPLACE INTO estado VALUES ('excel_casa', ?)",
                      (caminho,))
            c.execute("INSERT OR REPLACE INTO estado VALUES ('excel_casa_em', ?)",
                      (agora,))
    return rel


def texto_do_relatorio(rel):
    linhas = ["%d concursos no Excel%s" % (rel["total"],
                                            " (ENSAIO: nada foi gravado)" if rel["ensaio"] else "")]
    linhas.append("  ligados ao radar: %d (%d novos, %d à mão, %d pelo contrato no BASE)"
                  % (rel["ligadas"], rel["novas"], rel["manuais"], rel["pelo_base"]))
    if rel["aplicadas"]:
        linhas.append("  triagem: " + ", ".join(
            "%s %d" % (k, v) for k, v in sorted(rel["aplicadas"].items())))
        if list(rel["aplicadas"]) == ["guardado"]:
            linhas[-1] += " (só o registo; a triagem não se aplica sem --com-triagem)"
    if rel["lidos"]:
        linhas.append("  detalhes lidos ao DR para desempatar: %d" % rel["lidos"])
    if rel["em_lotes"]:
        linhas.append("  linhas em anúncios com lotes: %d, com o lote identificado: %d"
                      "%s"
                      % (rel["em_lotes"], rel["com_lote"],
                         (", pelo conjunto: %d" % rel["conjunto"]) if rel["conjunto"] else ""))
    linhas.append("  ambíguos: %d | sem correspondência: %d | sem anúncio no DR: %d"
                  " | fora do país: %d"
                  % (len(rel["ambiguas"]), len(rel["sem"]), rel["sem_dr"],
                     len(rel["fora"])))
    for ide, nome, ent, ano, cands in rel["ambiguas"]:
        linhas.append("    ? #%s %s — %s (%s): %s" % (ide, nome[:50], ent, ano,
                                                     ", ".join(cands)))
    for ide, nome, ent, ano in rel["sem"]:
        linhas.append("    - #%s %s — %s (%s)" % (ide, nome[:50], ent, ano))
    for nome, ref, st in rel["conflitos"]:
        linhas.append("    ! %s: o Excel diz %s, o radar decidiu outra coisa (%s)"
                      % (nome[:50], st, ref))
    return "\n".join(linhas)


# ------------------------------------------------------ para o painel

def registo_de(c, ref):
    """A linha da casa ligada a este anuncio, ou None."""
    r = c.execute("SELECT * FROM casa WHERE ref=? ORDER BY id LIMIT 1",
                  (ref,)).fetchone()
    return dict(r) if r else None


def ligar_a_mao(c, ide, ref, quem="Afonso", triagem=False, porque=""):
    """Liga uma linha do registo a um anuncio, resolvendo uma alteracao
    para o original. So escreve triagem com `triagem`. Devolve (ok,
    mensagem).

    `ref` pode ser "nenhum" -- nao ha anuncio no DR (consulta previa,
    ajuste directo, consulta preliminar, ou antes da base), com a razao
    em `porque` -- ou "?" para so anotar a razao ("nao sei") e deixar a
    linha por ligar."""
    import radar
    if not c.execute("SELECT 1 FROM casa WHERE id=?", (ide,)).fetchone():
        return False, "não há nenhum registo #%s" % ide
    if ref.strip().lower() in ("nenhum", "-", "nao", "não"):
        c.execute("UPDATE casa SET ref=NULL, ligacao='nenhum', candidatos='[]', "
                  "resultado='', aplicado_em=NULL, porque_sem_ref=? WHERE id=?",
                  (porque or "sem anúncio no DR", ide))
        return True, "#%s sem anúncio no DR (%s)" % (ide, porque or "sem razão")
    if ref.strip() == "?":
        c.execute("UPDATE casa SET porque_sem_ref=? WHERE id=?", (porque or "?", ide))
        return True, "#%s fica por ligar (%s)" % (ide, porque or "?")
    a = c.execute("SELECT ref, estado, altera FROM anuncios WHERE ref=?",
                  (ref,)).fetchone()
    if not a:
        return False, "não há nenhum anúncio %s" % ref
    if a["estado"] == "alteracao":
        raiz = radar.raiz_da_alteracao(c, ref, a["altera"] or "")
        if raiz:
            ref = raiz
    linha = c.execute("SELECT * FROM casa WHERE id=?", (ide,)).fetchone()
    if not linha:
        return False, "não há nenhum registo #%s" % ide
    d = dict(linha)
    for k in ("perfis", "concorrentes", "precos_perfis"):
        try:
            d[k] = json.loads(d.get(k) or "[]")
        except ValueError:
            d[k] = []
    resultado = aplicar(c, d, ref, quem) if triagem else "guardado"
    agora = datetime.now().strftime("%Y-%m-%d %H:%M")
    lotes = lotes_do_anuncio(c, ref)
    lote = lote_da_linha(d, lotes)
    c.execute("UPDATE casa SET ref=?, ligacao='manual', candidatos='[]', "
              "resultado=?, aplicado_em=?, porque_sem_ref=NULL, lote=? WHERE id=?",
              (ref, resultado, agora if resultado == "aplicado" else None, lote, ide))
    return True, "#%s ligado ao anúncio %s (%s%s)" % (
        ide, ref, resultado,
        (", lote %d de %d" % (lote, len(lotes))) if lote else
        (", %d lotes, o preço é o conjunto" % len(lotes)) if lote == 0 else
        (", %d lotes, lote por identificar" % len(lotes)) if lotes else "")


def desligar(c, ide):
    c.execute("UPDATE casa SET ref=NULL, ligacao='', resultado='', "
              "aplicado_em=NULL WHERE id=?", (ide,))


def desaplicar_da_copia(copia):
    """Desfaz o que uma importacao escreveu, repondo as PROPOSTAS tal
    como estao numa COPIA da base feita antes dela, e apaga do historico
    o que a importacao la escreveu. O registo (tabela casa) fica; as
    ligacoes ficam. Devolve (propostas repostas, linhas de historico
    apagadas).

    Existe porque a 02/09/2026 se importou e aplicou, e o Afonso decidiu
    a seguir que nada se aplica antes de o registo estar validado.

    Desde 15/09/2026 repoe propostas e nao colunas do anuncio -- e por
    isso repor tambem sabe APAGAR: uma proposta que a importacao criou
    do nada nao estava na copia, e deixa-la la era a importacao ficar
    meia desfeita. Uma copia de ANTES da escada nao tem a tabela, e ai
    nao se desfaz nada: di-lo devolvendo zero, em vez de apagar tudo o
    que encontrar."""
    import radar
    import sqlite3
    antes = sqlite3.connect("file:%s?mode=ro" % copia.replace("\\", "/"), uri=True)
    antes.row_factory = sqlite3.Row
    repostos = apagadas = 0
    with radar.liga() as c:
        refs = [r["ref"] for r in c.execute(
            "SELECT DISTINCT ref FROM casa WHERE ref IS NOT NULL "
            "AND resultado IN ('aplicado', 'conflito', 'igual')")]
        try:
            antes.execute("SELECT 1 FROM propostas LIMIT 1")
        except sqlite3.OperationalError:
            refs = []
        for ref in refs:
            velhas = antes.execute(
                "SELECT %s FROM propostas WHERE ref=?"
                % ", ".join(radar.COLUNAS_DA_PROPOSTA), (ref,)).fetchall()
            c.execute("DELETE FROM propostas WHERE ref=?", (ref,))
            for v in velhas:
                c.execute("INSERT INTO propostas (%s) VALUES (%s)"
                          % (", ".join(radar.COLUNAS_DA_PROPOSTA),
                             ", ".join("?" * len(radar.COLUNAS_DA_PROPOSTA))),
                          [v[k] for k in radar.COLUNAS_DA_PROPOSTA])
            repostos += 1
        apagadas = c.execute("DELETE FROM historico WHERE quem='Excel'").rowcount
        c.execute("UPDATE casa SET resultado='guardado', aplicado_em=NULL "
                  "WHERE ref IS NOT NULL AND resultado != 'fora'")
    antes.close()
    return repostos, apagadas


# ------------------------------------------ o modelo da casa (8/09/2026)
#
# Decisao do Afonso a 8/09/2026: em vez de o radar tentar perceber o
# Excel antigo (Analise_Concursos_Publicos.xlsm, feito para outra coisa,
# com folhas C_ e consolidadores VBA), **o radar dita o modelo**: um
# .xlsx gerado aqui, com as colunas que a aplicacao precisa e listas de
# escolha onde ha vocabulario, que o utilizador preenche e carrega em
# Configuracoes > Importar dados, com ensaio antes de gravar. O leitor
# do Excel antigo (ler_excel, importar, ligar_a_mao) fica acima, sem
# comando que o chame: e historico, e os testes dele continuam a valer.
#
# A chave e a REFERENCIA DO ANUNCIO no DR ("1947/2026"), que a ficha
# mostra: liga sem adivinhar, e uma linha sem anuncio e um erro que se
# ve no ensaio, nao um palpite.

COLUNAS_MODELO = (
    ("Referência do anúncio", "ref"),
    ("Lote", "lote"),
    ("Estado", "status"),
    ("Razão de não participação", "razao"),
    ("Valor da proposta (€)", "valor_proposta"),
    ("Lugar", "lugar"),
    ("Concorrentes (separados por ;)", "concorrentes"),
    ("Responsável", "responsavel"),
    ("Notas", "notas"),
)
ESTADOS_MODELO = ("Não fomos", "Submetido", "Ganho", "Perdido")
FOLHA_MODELO = "Registo"
PASTA_IMPORTACOES = "importacoes"
RX_REF = re.compile(r"^\s*(\d{1,6})\s*[/\-\s]\s*(\d{4})\s*$")


def escrever_modelo(caminho):
    """Gera o .xlsx vazio: a folha Registo com os cabecalhos e as listas
    de escolha (Estado, Razao) ate a linha 500, e uma folha de
    instrucoes com um exemplo. Devolve o caminho."""
    import radar
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
    wb = Workbook()
    ws = wb.active
    ws.title = FOLHA_MODELO
    larguras = (22, 8, 14, 28, 20, 8, 44, 18, 40)
    for i, ((titulo, _), largura) in enumerate(zip(COLUNAS_MODELO, larguras), 1):
        c = ws.cell(row=1, column=i, value=titulo)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="17557F")
        c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = largura
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"
    dv_estado = DataValidation(type="list", formula1='"%s"' % ",".join(ESTADOS_MODELO),
                               allow_blank=True, showErrorMessage=True,
                               errorTitle="Estado", error="Escolhe um da lista.")
    dv_razao = DataValidation(type="list",
                              formula1='"%s"' % ",".join(radar.MOTIVOS_ABANDONO),
                              allow_blank=True, showErrorMessage=False)
    ws.add_data_validation(dv_estado)
    ws.add_data_validation(dv_razao)
    dv_estado.add("C2:C500")
    dv_razao.add("D2:D500")
    inst = wb.create_sheet("Instruções")
    inst.column_dimensions["A"].width = 110
    linhas = [
        "Como preencher a folha «Registo» — uma linha por concurso, ou por lote quando o concurso tem lotes.",
        "",
        "Referência do anúncio: a referência do DR tal como a ficha do radar a mostra, ex. 1947/2026. É obrigatória e é o que liga a linha ao anúncio.",
        "Lote: o número do lote (1, 2, 3…) quando o concurso tem lotes e a linha é de um lote. Vazio quando não há lotes ou quando se foi ao conjunto.",
        "Estado: um da lista — Não fomos, Submetido, Ganho, Perdido.",
        "Razão de não participação: só quando o estado é «Não fomos» — %s (ou outra, em texto livre)." % ", ".join(radar.MOTIVOS_ABANDONO),
        "Valor da proposta (€): o que propusemos, em número (ex. 54432 ou 54432,50). Vazio se não fomos.",
        "Lugar: a posição no relatório preliminar (1, 2, 3…). Vazio se ainda não há relatório.",
        "Concorrentes: os nomes separados por ponto e vírgula, por ordem de classificação, ex. Empresa A; Empresa B; Empresa C.",
        "Responsável: quem da casa acompanha este concurso (nome).",
        "Notas: texto livre.",
        "",
        "Exemplo:  1947/2026 | 2 | Ganho |  | 169344 | 1 | Nós; Empresa B; Empresa C | Afonso | contrato de 24 meses",
        "",
        "Depois de preencher, carrega o ficheiro em Configurações › Importar dados. O radar mostra um ensaio (o que liga a que anúncio, o que não liga e porquê) e só grava quando confirmares.",
        "Uma linha repetida (mesma referência e mesmo lote) substitui a anterior. Linhas com erro não entram; as outras entram.",
    ]
    for i, t in enumerate(linhas, 1):
        inst.cell(row=i, column=1, value=t).alignment = Alignment(wrap_text=True, vertical="top")
    inst.cell(row=1, column=1).font = Font(bold=True)
    wb.save(caminho)
    return caminho


def ref_limpa(texto):
    """"1947/2026", "1947-2026", " 1947 / 2026 " -> "1947/2026"; senao ''."""
    m = RX_REF.match(str(texto or ""))
    return "%s/%s" % (int(m.group(1)), m.group(2)) if m else ""


def _celula(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def ler_modelo(caminho):
    """Le o .xlsx preenchido: [{linha, ref, lote, status, razao,
    valor_proposta, lugar, concorrentes, responsavel, notas, erros}].

    Nao decide nada sobre a base -- isso e o ensaio_modelo(). Aqui so
    se normaliza e se apontam os erros de forma: referencia ilegivel,
    estado fora da lista, lote ou lugar que nao sao numeros."""
    from openpyxl import load_workbook
    wb = load_workbook(caminho, read_only=True, data_only=True)
    ws = wb[FOLHA_MODELO] if FOLHA_MODELO in wb.sheetnames else wb.active
    chaves = [c for _, c in COLUNAS_MODELO]
    saida = []
    for n, valores in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
        valores = list(valores or ())[:len(chaves)]
        valores += [None] * (len(chaves) - len(valores))
        bruto = dict(zip(chaves, valores))
        if not any(_celula(v) for v in valores):
            continue
        erros = []
        linha = {"linha": n, "erros": erros}
        linha["ref"] = ref_limpa(bruto["ref"])
        if not linha["ref"]:
            erros.append("referência do anúncio ilegível (ex.: 1947/2026)")
        lote_txt = _celula(bruto["lote"])
        try:
            linha["lote"] = int(float(lote_txt.replace(",", "."))) if lote_txt else None
            if linha["lote"] is not None and linha["lote"] < 1:
                erros.append("o lote tem de ser 1 ou mais")
        except ValueError:
            linha["lote"] = None
            erros.append("lote «%s» não é um número" % lote_txt)
        estado = _celula(bruto["status"])
        por_norma = {_norma(e): e for e in ESTADOS_MODELO}
        linha["status"] = por_norma.get(_norma(estado), "")
        if not linha["status"]:
            erros.append("estado «%s» não está na lista (%s)"
                         % (estado, ", ".join(ESTADOS_MODELO)))
        razao = _celula(bruto["razao"])
        linha["razao"] = MAPA_RAZAO.get(_norma(razao), razao) if razao else ""
        linha["valor_proposta"] = _num(bruto["valor_proposta"])
        lugar_txt = _celula(bruto["lugar"])
        try:
            linha["lugar"] = int(float(lugar_txt.replace(",", "."))) if lugar_txt else None
        except ValueError:
            linha["lugar"] = None
            erros.append("lugar «%s» não é um número" % lugar_txt)
        nomes = [p.strip() for p in re.split(r"[;\n]", _celula(bruto["concorrentes"])) if p.strip()]
        linha["concorrentes"] = [{"lugar": i, "nome": nome} for i, nome in enumerate(nomes, 1)]
        linha["responsavel"] = _celula(bruto["responsavel"])[:60]
        linha["notas"] = _celula(bruto["notas"])[:2000]
        linha["nome"] = ""
        saida.append(linha)
    wb.close()
    return saida


def ensaio_modelo(c, linhas):
    """Cruza as linhas lidas com a base, sem gravar: poe em cada uma o
    titulo do anuncio, a lista de problemas e `ok`. Devolve (linhas,
    contagens)."""
    vistas = {}
    for l in linhas:
        problemas = list(l["erros"])
        l["titulo"] = ""
        a = None
        if l["ref"]:
            a = c.execute("SELECT ref, titulo, estado, lotes FROM anuncios WHERE ref=?",
                          (l["ref"],)).fetchone()
            if not a:
                problemas.append("não há anúncio %s na base" % l["ref"])
            else:
                l["titulo"] = a["titulo"] or ""
                if a["estado"] == "alteracao":
                    problemas.append("%s é uma republicação; usa a referência do anúncio original" % l["ref"])
                lotes = lotes_do_anuncio(c, l["ref"])
                if l["lote"] is not None and not lotes:
                    problemas.append("o anúncio não declara lotes; deixa o lote vazio")
                elif l["lote"] is not None and l["lote"] not in [x["n"] for x in lotes]:
                    problemas.append("o anúncio tem %d lotes e não tem o lote %d"
                                     % (len(lotes), l["lote"]))
                elif l["lote"] is None and lotes:
                    l.setdefault("avisos", []).append(
                        "o anúncio tem %d lotes; sem lote, a linha conta como o conjunto" % len(lotes))
        chave = (l["ref"], l["lote"])
        if l["ref"] and chave in vistas:
            problemas.append("repete a linha %d (mesma referência e lote)" % vistas[chave])
        elif l["ref"]:
            vistas[chave] = l["linha"]
        if _norma(l.get("status")) == "nao fomos" and (l.get("valor_proposta") or l.get("lugar")):
            l.setdefault("avisos", []).append("«Não fomos» com proposta ou lugar: ficam guardados, mas não contam")
        l["problemas"] = problemas
        l["ok"] = not problemas
    contagens = {"total": len(linhas), "ok": sum(1 for l in linhas if l["ok"]),
                 "com_erro": sum(1 for l in linhas if not l["ok"]),
                 "anuncios": len({l["ref"] for l in linhas if l["ok"]})}
    return linhas, contagens


# A prioridade quando um anuncio tem varias linhas (lotes) com estados
# diferentes: o anuncio fica no melhor deles -- ganhamos um lote, o
# cartao esta no Ganho, e a separacao no fim mostra os perdidos.
_PRIORIDADE = {"ganho": 0, "submetido": 1, "perdido": 2, "nao fomos": 3}


def aplicar_modelo(c, linhas, quem="modelo"):
    """Grava as linhas `ok` na tabela casa (folha='modelo') e escreve a
    triagem nos anuncios. Devolve {"gravadas", "aplicadas", "anuncios",
    "resultados": {ref: resultado de aplicar()}}."""
    import radar
    agora = datetime.now().strftime("%Y-%m-%d %H:%M")
    por_ref = {}
    gravadas = 0
    for l in linhas:
        if not l.get("ok"):
            continue
        c.execute("DELETE FROM casa WHERE ref=? AND folha='modelo' AND "
                  "COALESCE(lote,0)=COALESCE(?,0)", (l["ref"], l["lote"] if l["lote"] is not None else 0))
        c.execute(
            "INSERT INTO casa (nome, status, razao, valor_proposta, lugar, notas, folha, "
            "concorrentes, ref, ligacao, candidatos, fora, resultado, importado_em, "
            "aplicado_em, lote) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
            (l.get("titulo") or "", l["status"], l["razao"] or None, l["valor_proposta"],
             l["lugar"], l["notas"] or None, "modelo",
             json.dumps(l["concorrentes"], ensure_ascii=False), l["ref"], "ref", "[]",
             0, "aplicado", agora, agora, l["lote"] if l["lote"] is not None else 0))
        gravadas += 1
        por_ref.setdefault(l["ref"], []).append(l)
    resultados = {}
    for ref, grupo in por_ref.items():
        melhor = sorted(grupo, key=lambda l: _PRIORIDADE.get(_norma(l["status"]), 9))[0]
        resultados[ref] = aplicar(c, melhor, ref, quem=quem)
        responsavel = next((l["responsavel"] for l in grupo if l.get("responsavel")), "")
        if responsavel:
            c.execute("INSERT OR IGNORE INTO pessoas (nome) VALUES (?)", (responsavel,))
            # Na PROPOSTA, e nao no anuncio (15/09/2026): quem trata de um
            # concurso e quem trata da proposta, e um anuncio por ver nao
            # tem dono porque ainda nao ha nada para tratar.
            c.execute("UPDATE propostas SET responsavel=? WHERE ref=?",
                      (responsavel, ref))
    return {"gravadas": gravadas, "anuncios": len(por_ref),
            "aplicadas": sum(1 for r in resultados.values() if r == "aplicado"),
            "resultados": resultados}
